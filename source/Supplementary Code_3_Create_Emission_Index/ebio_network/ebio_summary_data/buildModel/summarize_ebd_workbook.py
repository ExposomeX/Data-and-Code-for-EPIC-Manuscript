"""Summarize the supplied E-BIO network workbook without altering its records.

Run with the project Python environment. Outputs remain beside this script.
"""

from __future__ import annotations

import csv
import json
import re
from collections import Counter
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook


SOURCE = Path(r"D:\emi_association\稿件\Supplementary Data\SM\Table S. EBD network egdes and nodes data.xlsx")
OUT = Path(__file__).resolve().parent
GRAPHS = ("EGoD", "EPPD")


def write_csv(name: str, rows: list[dict]) -> None:
    if not rows:
        return
    keys = list(rows[0])
    with (OUT / name).open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=keys)
        writer.writeheader()
        writer.writerows(rows)


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def read_nodes(sheet) -> tuple[dict, dict]:
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    idx = {str(item): position for position, item in enumerate(header)}
    for key in ("EXID", "Group name", "Degree"):
        if key not in idx:
            raise ValueError(f"{sheet.title}: missing {key}")
    identifiers: dict[str, str] = {}
    group_counts = Counter()
    missing_id = duplicate_id = missing_group = 0
    degrees = []
    for row in rows:
        if all(item is None for item in row):
            continue
        ident = text(row[idx["EXID"]])
        group = text(row[idx["Group name"]]) or "[missing]"
        if not ident:
            missing_id += 1
        elif ident in identifiers:
            duplicate_id += 1
        else:
            identifiers[ident] = group
        if group == "[missing]":
            missing_group += 1
        group_counts[group] += 1
        raw_degree = row[idx["Degree"]]
        if isinstance(raw_degree, (int, float)) and np.isfinite(raw_degree):
            degrees.append(float(raw_degree))
    n = sum(group_counts.values())
    info = dict(
        node_rows=n,
        unique_node_ids=len(identifiers),
        missing_node_ids=missing_id,
        duplicated_node_ids=duplicate_id,
        missing_node_groups=missing_group,
        degree_median=float(np.median(degrees)) if degrees else None,
        degree_max=float(np.max(degrees)) if degrees else None,
        node_group_counts=dict(sorted(group_counts.items())),
    )
    return info, identifiers


def read_edges(sheet, node_ids: set[str]) -> dict:
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    idx = {str(item): position for position, item in enumerate(header)}
    required = ("Source", "Target", "Source group", "Target group", "Interaction", "Database", "Edge type")
    for key in required:
        if key not in idx:
            raise ValueError(f"{sheet.title}: missing {key}")
    transition = Counter()
    database = Counter()
    database_family = Counter()
    interaction = Counter()
    edge_type = Counter()
    unique_pairs: set[tuple[str, str]] = set()
    unique_records: set[tuple[str, str, str, str, str]] = set()
    source_ids, target_ids, missing_nodes = set(), set(), set()
    missing_source = missing_target = self_loops = 0
    edge_rows = 0
    for row in rows:
        if all(item is None for item in row):
            continue
        edge_rows += 1
        source, target = text(row[idx["Source"]]), text(row[idx["Target"]])
        sg, tg = text(row[idx["Source group"]]), text(row[idx["Target group"]])
        rel, db, etype = (text(row[idx[key]]) for key in ("Interaction", "Database", "Edge type"))
        transition[(sg or "[missing]", tg or "[missing]")] += 1
        database[db or "[missing]"] += 1
        if db.startswith("PMID:"):
            family = "PMID-labelled"
        elif db.startswith("ctd^"):
            family = "ctd^*"
        elif db.startswith("ctd"):
            family = "ctd"
        else:
            family = re.split(r"[\^|]", db, maxsplit=1)[0] or "[missing]"
        database_family[family] += 1
        interaction[rel or "[missing]"] += 1
        edge_type[etype or "[missing]"] += 1
        if not source:
            missing_source += 1
        else:
            source_ids.add(source)
            if source not in node_ids:
                missing_nodes.add(source)
        if not target:
            missing_target += 1
        else:
            target_ids.add(target)
            if target not in node_ids:
                missing_nodes.add(target)
        if source and target:
            unique_pairs.add((source, target))
            unique_records.add((source, target, rel, db, etype))
            if source == target:
                self_loops += 1
    return dict(
        edge_rows=edge_rows,
        unique_directed_pairs=len(unique_pairs),
        unique_edge_records=len(unique_records),
        duplicate_directed_pair_rows=edge_rows - len(unique_pairs),
        duplicate_full_record_rows=edge_rows - len(unique_records),
        unique_source_ids=len(source_ids),
        unique_target_ids=len(target_ids),
        unique_endpoint_ids=len(source_ids | target_ids),
        endpoint_ids_absent_from_node_sheet=len(missing_nodes),
        missing_source_rows=missing_source,
        missing_target_rows=missing_target,
        self_loop_rows=self_loops,
        transition_counts={f"{a} → {b}": n for (a, b), n in sorted(transition.items())},
        database_counts=dict(sorted(database.items())),
        database_family_counts=dict(sorted(database_family.items())),
        interaction_counts=dict(sorted(interaction.items())),
        edge_type_counts=dict(sorted(edge_type.items())),
        source_id_set=source_ids,
        target_id_set=target_ids,
        endpoint_id_set=source_ids | target_ids,
        node_id_set=node_ids,
    )


def draw_transition_figure(summaries: dict) -> None:
    plt.rcParams.update(
        {"font.family": "sans-serif", "font.sans-serif": ["Arial", "DejaVu Sans"],
         "font.size": 8, "pdf.fonttype": 42,
         "svg.fonttype": "none", "axes.spines.top": False,
         "axes.spines.right": False, "axes.linewidth": 0.8}
    )
    rows = []
    for graph in GRAPHS:
        for transition, count in summaries[graph]["transition_counts"].items():
            rows.append((graph, transition, count))
    rows.sort(key=lambda item: (GRAPHS.index(item[0]), -item[2]))
    labels = [f"{graph}  |  {transition}" for graph, transition, _ in rows]
    values = [count for _, _, count in rows]
    if not values or min(values) <= 0:
        raise ValueError("Log-scale network composition requires positive edge counts")
    colors = ["#39749b" if graph == "EGoD" else "#b9704f" for graph, _, _ in rows]
    fig, ax = plt.subplots(figsize=(7.2, max(3.0, 0.32 * len(rows) + 1.2)))
    ypos = np.arange(len(rows))
    ax.barh(ypos, values, color=colors, height=0.67)
    ax.set_yticks(ypos, labels)
    ax.invert_yaxis()
    ax.set_xlabel("Edge records (logarithmic scale)")
    ax.set_xscale("log")
    ax.grid(False)
    ax.set_xlim(30, max(values) * 1.7)
    for y, value in zip(ypos, values):
        ax.text(value * 1.4, y, f"{value:,}", va="center", fontsize=7)
    fig.tight_layout(pad=1.2)
    stem = OUT / "Fig_EBD_network_edge_composition"
    fig.savefig(stem.with_suffix(".pdf"), bbox_inches="tight")
    fig.savefig(stem.with_suffix(".svg"), bbox_inches="tight")
    fig.savefig(stem.with_suffix(".png"), dpi=600, bbox_inches="tight")
    fig.savefig(stem.with_suffix(".tiff"), dpi=600, bbox_inches="tight")
    plt.close(fig)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    book = load_workbook(SOURCE, read_only=True, data_only=True)
    if set(book.sheetnames) != {f"{g}_{kind}" for g in GRAPHS for kind in ("nodes", "edges")}:
        raise ValueError(f"Unexpected workbook sheets: {book.sheetnames}")
    results = {}
    groups, transitions, databases, families, relations = [], [], [], [], []
    graph_nodes = {}
    for graph in GRAPHS:
        n, nodes = read_nodes(book[f"{graph}_nodes"])
        graph_nodes[graph] = nodes
        e = read_edges(book[f"{graph}_edges"], set(nodes))
        e["node_ids_not_used_by_edges"] = len(set(nodes) - e["endpoint_id_set"])
        for key in ("source_id_set", "target_id_set", "endpoint_id_set", "node_id_set"):
            e.pop(key)
        results[graph] = {**n, **e}
        groups += [dict(network=graph, group=key, count=value) for key, value in n["node_group_counts"].items()]
        transitions += [dict(network=graph, transition=key, count=value) for key, value in e["transition_counts"].items()]
        databases += [dict(network=graph, database=key, count=value) for key, value in e["database_counts"].items()]
        families += [dict(network=graph, source_label_family=key, count=value) for key, value in e["database_family_counts"].items()]
        relations += [dict(network=graph, interaction=key, count=value) for key, value in e["interaction_counts"].items()]
        print(graph, "node rows", n["node_rows"], "edge rows", e["edge_rows"], flush=True)
    book.close()
    write_csv("Table_network_overview.csv", [{"network": graph, **{k: v for k, v in results[graph].items() if not isinstance(v, dict)}} for graph in GRAPHS])
    write_csv("Table_node_groups.csv", groups)
    write_csv("Table_edge_transitions.csv", transitions)
    write_csv("Table_database_labels.csv", databases)
    write_csv("Table_database_label_families.csv", families)
    write_csv("Table_interaction_labels.csv", relations)
    overlap = {}
    for group in ("Exposure", "Disease"):
        left = {key for key, value in graph_nodes["EGoD"].items() if value == group}
        right = {key for key, value in graph_nodes["EPPD"].items() if value == group}
        overlap[group] = dict(EGoD=len(left), EPPD=len(right), shared=len(left & right), union=len(left | right))
    write_csv("Table_cross_network_overlap.csv", [dict(node_group=key, **value) for key, value in overlap.items()])
    compact = {
        graph: {key: value for key, value in results[graph].items() if key not in ("database_counts", "interaction_counts")}
        for graph in GRAPHS
    }
    compact["cross_network_overlap"] = overlap
    (OUT / "network_summary.json").write_text(json.dumps(compact, indent=2, ensure_ascii=False), encoding="utf-8")
    draw_transition_figure(results)
    print(json.dumps(compact, indent=2, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
